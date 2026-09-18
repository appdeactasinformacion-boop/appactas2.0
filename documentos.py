# ==============================================================
# PLANTILLAS WORD: lectura de campos y generación del documento
# ==============================================================

import re

import docx
import streamlit as st
from docxtpl import DocxTemplate


def get_fields_from_template(template_path):
    doc = docx.Document(template_path)
    found_fields = set()
    pattern = re.compile(r'\{\{.*?\}\}|\{%.*?%\}')
    for para in doc.paragraphs:
        found_fields.update(pattern.findall(para.text))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    found_fields.update(pattern.findall(para.text))
    return list(found_fields)


def normalizar_listas(data):
    # La IA a veces devuelve la clave "responsablen" y las plantillas usan
    # "resposablen": se aceptan ambas escrituras para que nunca salga vacío.
    lista_comp = data.get("COMPROMISOS_DE_REUNION", [])
    if isinstance(lista_comp, list):
        for item in lista_comp:
            if isinstance(item, dict) and not item.get("resposablen") and item.get("responsablen"):
                item["resposablen"] = item["responsablen"]
    claves = {
        "ASISTENTES_REUNION": ["nombreasistentereu", "cargoasistentereunion"],
        "TEMAS_TRATADOS_N": ["tema", "responsablet"],
        "COMPROMISOS_DE_REUNION": ["compromiso", "resposablen", "fechac", "fechas"],
        "TEMAS_TRATADOS": ["tema", "desarrollo"],
        "COMPROMISOS_R": ["compromiso", "responsable", "fechaejecucion"],
    }
    for clave, campos in claves.items():
        lista = data.get(clave, [])
        if not isinstance(lista, list):
            lista = []
        for item in lista:
            for campo in campos:
                item.setdefault(campo, "N/A")
        data[clave.lower()] = lista
        data.pop(clave, None)


def create_word_document(template_path, data, elaborada_por="N/A", cargo="N/A"):
    try:
        doc = DocxTemplate(template_path)
        normalizar_listas(data)
        data["ACTA_ELABORADA_POR"] = elaborada_por or "N/A"
        data["CARGO_ELA"] = cargo or "N/A"
        doc.render(data)
        output_path = "acta_generada.docx"
        doc.save(output_path)
        return output_path
    except Exception as e:
        st.error(f"No se pudo generar el documento: {e}")
        return None
